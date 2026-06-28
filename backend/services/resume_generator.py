from datetime import date
from pathlib import Path
import re
import shutil
import subprocess

from docx import Document

from services.file_service import active_resume_source, master_text
from services.resume_agents import run_agent_pipeline

BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "storage" / "generated_resumes"
DRAFT_DIR = BASE_DIR / "storage" / "resume_drafts"


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_") or "ManualJD"


def unique_path(company: str, role: str) -> Path:
    stem = f"Local_Candidate_{safe_name(company)}_{safe_name(role)}_{date.today():%Y%m%d}"
    path = OUTPUT_DIR / f"{stem}.docx"
    counter = 1
    while path.exists():
        path = OUTPUT_DIR / f"{stem}_resume_{counter}.docx"
        counter += 1
    return path


def unique_draft_path(company: str, role: str) -> Path:
    DRAFT_DIR.mkdir(parents=True, exist_ok=True)
    stem = f"Draft_{safe_name(company)}_{safe_name(role)}_{date.today():%Y%m%d}"
    path = DRAFT_DIR / f"{stem}.docx"
    counter = 1
    while path.exists():
        path = DRAFT_DIR / f"{stem}_{counter}.docx"
        counter += 1
    return path


def _txt_template(source: Path) -> Path:
    converted = source.parent / "editable_base.docx"
    if converted.exists() and converted.stat().st_mtime >= source.stat().st_mtime:
        return converted

    document = Document()
    for block in source.read_text(encoding="utf-8", errors="ignore").splitlines():
        text = block.strip()
        if not text:
            continue
        upper = text.upper()
        if upper in {"SUMMARY", "SKILLS", "EXPERIENCE"}:
            document.add_paragraph(f"{upper}:")
        elif text.startswith("- "):
            document.add_paragraph(text[2:], style="List Paragraph")
        else:
            document.add_paragraph(text)
    document.save(converted)
    return converted


def _editable_template(source: Path) -> Path:
    if source.suffix.lower() == ".docx":
        return source
    if source.suffix.lower() == ".txt":
        return _txt_template(source)
    if source.suffix.lower() == ".pdf":
        # pdf2docx brings in a sizeable document-conversion stack.  Loading it only
        # for PDF input keeps the local API responsive for the normal DOCX/TXT path.
        from pdf2docx import Converter

        converted = source.parent / "editable_base.docx"
        if not converted.exists() or converted.stat().st_mtime < source.stat().st_mtime:
            converter = Converter(str(source))
            try:
                converter.convert(str(converted))
            finally:
                converter.close()
        return converted
    raise ValueError(
        "Resume generation requires a DOCX, PDF, or TXT base resume."
    )


def create_resume_draft(
    company: str,
    role: str,
    jd_text: str,
    analysis: dict,
) -> tuple[Path, dict]:
    source_text = master_text()
    source_path = active_resume_source()
    if not source_text or source_path is None:
        raise ValueError("Upload and select a main resume before generating a tailored version.")

    template = _editable_template(source_path)
    path = unique_draft_path(company, role)
    shutil.copy2(template, path)
    document = Document(path)
    base_document = Document(template)

    report = run_agent_pipeline(
        document=document,
        base_document=base_document,
        source_text=source_text,
        role=role,
        jd_text=jd_text,
        analysis=analysis,
    )
    document.core_properties.title = f"{role} - {company}"
    document.core_properties.subject = (
        "Tailored and independently verified against the selected base resume"
    )
    document.save(path)
    return path, report


def approve_resume_draft(draft_path: str | Path, company: str, role: str) -> Path:
    source = Path(draft_path)
    if not source.exists():
        raise ValueError("The saved resume draft could not be found.")
    path = unique_path(company, role)
    shutil.copy2(source, path)
    return path


def pdf_for_resume(docx_path: str | Path) -> Path:
    """Create a local PDF beside a final DOCX without overwriting another version."""
    source = Path(docx_path)
    output_dir = source.parent
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(source)],
            check=True,
            capture_output=True,
            text=True,
            timeout=60,
        )
    except FileNotFoundError as exc:
        raise ValueError("PDF output requires LibreOffice (soffice) to be installed locally.") from exc
    except subprocess.CalledProcessError as exc:
        raise ValueError(f"LibreOffice could not convert the approved resume to PDF: {exc.stderr.strip()}") from exc
    path = output_dir / f"{source.stem}.pdf"
    if not path.exists():
        raise ValueError("LibreOffice did not produce the expected PDF output.")
    return path


def generate_resume(
    company: str,
    role: str,
    jd_text: str,
    analysis: dict,
) -> tuple[Path, dict]:
    """Legacy convenience entry point. New UI flows must approve a draft first."""
    draft, report = create_resume_draft(company, role, jd_text, analysis)
    return approve_resume_draft(draft, company, role), report
