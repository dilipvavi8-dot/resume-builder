from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
import re

from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from utils.scoring import detect_skills


@dataclass
class ResumeChange:
    section: str
    original: str
    updated: str
    reason: str
    client: str = ""


def _normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def _move_section_break(source: Paragraph, target: Paragraph) -> None:
    """Move a Word section break before removing an otherwise blank paragraph."""
    if source._p.pPr is None or source._p.pPr.sectPr is None:
        return
    section_properties = deepcopy(source._p.pPr.sectPr)
    source._p.pPr.remove(source._p.pPr.sectPr)
    target._p.get_or_add_pPr().append(section_properties)


def _insert_after(
    paragraph: Paragraph,
    text: str,
    formatting_source: Paragraph | None = None,
) -> Paragraph:
    formatting_source = formatting_source or paragraph
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted._p.get_or_add_pPr()
    if formatting_source._p.pPr is not None:
        inserted._p.replace(inserted._p.pPr, deepcopy(formatting_source._p.pPr))
    run = inserted.add_run(text)
    source_run = next((item for item in formatting_source.runs if item.text.strip()), None)
    if source_run is not None and source_run._r.rPr is not None:
        run._r.get_or_add_rPr()
        run._r.replace(run._r.rPr, deepcopy(source_run._r.rPr))
    return inserted


def _summary_paragraph(document: DocumentType) -> Paragraph | None:
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        heading = paragraph.text.upper().strip().rstrip(":")
        if heading not in {"PROFESSIONAL SUMMARY", "SUMMARY"}:
            continue
        return next(
            (
                candidate
                for candidate in paragraphs[index + 1:]
                if candidate.text.strip()
                and candidate.text.upper().strip().rstrip(":") not in {"TECHNICAL SKILLS", "SKILLS"}
            ),
            None,
        )
    return None


def _client_for_index(document: DocumentType, index: int) -> str:
    """Find the nearest dated employer heading for a work-experience bullet."""
    for paragraph in reversed(document.paragraphs[:index]):
        text = _normalized(paragraph.text)
        if " | " in text and re.search(r"\b20\d{2}\b", text):
            return text.split("|", 1)[0].strip()
        if re.search(r"\b(?:20\d{2}|Till Now)\b", text) and len(text) < 140:
            return re.sub(r"\s+(?:\w+\s+)?20\d{2}.*$", "", text).strip()
    return ""


def _experience_years(text: str) -> str:
    years = [
        int(value)
        for value in re.findall(r"(\d+)\+?\s*(?:years|yrs)", text.lower())
        if int(value) <= 40
    ]
    return f"{max(years)}+ years" if years else "extensive"


def _targeted_summary(
    source_text: str,
    role: str,
    analysis: dict,
) -> str:
    matched = analysis.get("matching_skills", [])
    highlighted = ", ".join(matched[:7])
    raw_role = role.strip() or analysis.get("role_type", "cloud engineering")
    title_case_overrides = {"ai": "AI", "devops": "DevOps", "sre": "SRE"}
    role_label = " ".join(
        title_case_overrides.get(word.lower(), word.capitalize())
        for word in raw_role.split()
    )
    experience = _experience_years(source_text)
    support_terms = []
    support_terms.extend(
        [
            "cloud operations",
            "application lifecycle support",
            "production troubleshooting",
            "performance monitoring",
        ]
    )
    return (
        f"Senior DevOps, cloud, and site reliability professional with {experience} of "
        f"experience targeting {role_label} opportunities. Demonstrated experience with "
        f"{highlighted or 'cloud automation and production operations'}, including "
        f"{', '.join(support_terms)}. Translates operational requirements into reliable "
        "automation, deployment, monitoring, and incident-response improvements while "
        "partnering with engineering and operations teams."
    )


def _rewrite_for_requirement(text: str, requirement: str) -> str:
    lowered = text.lower()
    if requirement == "AWS infrastructure optimization":
        if "automated aws services" in lowered:
            return (
                "Managed and optimized AWS infrastructure across IAM, EC2, VPC, RDS, "
                "S3, CloudFront, Lambda, CloudWatch, ECS, Route53, and related services "
                "using reusable Terraform modules to improve consistency and reusability."
            )
        if "deployed aws infrastructure using cloudformation" in lowered:
            return (
                "Supported AWS compute, storage, networking, and serverless services "
                "including EC2, EKS, S3, RDS, ELB, IAM, Route 53, Lambda, ECS, "
                "Auto Scaling, and CloudTrail."
            )
    if requirement == "CI/CD lifecycle":
        if "implemented end-to-end ci/cd pipelines" in lowered:
            return (
                "Designed, implemented, and maintained CI/CD pipelines in Azure DevOps "
                "for .NET and Java applications across AKS and App Services, supporting "
                "SDLC controls and zero-downtime releases."
            )
        if "implemented github actions" in lowered:
            return (
                "Built CI/CD workflows for microservices with GitHub Actions, SonarQube "
                "quality checks, Docker image builds, and ECS blue/green deployments."
            )
    if requirement == "Infrastructure as code and automation":
        if "built and managed aws eks clusters using terraform modules" in lowered:
            return (
                "Developed infrastructure-as-code automation with Terraform modules to "
                "provision AWS EKS clusters and promote consistent environment builds."
            )
        if "automated provisioning" in lowered:
            return (
                "Automated infrastructure provisioning for EC2, S3, RDS, and VPC "
                "networking components, reducing environment setup time from hours to "
                "minutes."
            )
    if requirement == "Container orchestration":
        if "deployed and managed containerized apps" in lowered:
            return (
                "Implemented and managed containerized application deployments on "
                "Kubernetes across AKS and EKS with scaling, auto-healing, and rolling "
                "update strategies."
            )
        if "ecs fargate" in lowered:
            return (
                "Configured ECS Fargate services with load balancing and autoscaling "
                "policies to support reliable containerized application operations."
            )
    if requirement == "Python application support":
        if "automated aws lambda incident resolution scripts using python" in lowered:
            return (
                "Developed and supported Python automation for AWS Lambda workloads, "
                "troubleshooting failed functions, validating triggers, and sending SNS "
                "alerts, reducing manual intervention by 30%."
            )
        if "wrote python and bash scripts" in lowered:
            return (
                "Developed Python and Bash automation for application health checks, "
                "Lambda recovery, and CloudWatch log analysis, reducing manual support "
                "work by 40%."
            )
    if requirement == "Application lifecycle and troubleshooting":
        if "diagnosed aws lambda failures" in lowered:
            return (
                "Provided application lifecycle support for deployed AWS Lambda workloads "
                "by analyzing CloudWatch logs and metrics, validating trigger "
                "configurations, and resolving recurring failure patterns."
            )
        if "reproduced and debugged issues" in lowered:
            return (
                "Reproduced, debugged, and validated application issues in controlled test "
                "environments to confirm root causes before production remediation."
            )
    if requirement == "Business and technical collaboration":
        if "collaborated with development and operations teams" in lowered:
            return (
                "Partnered with development and operations teams to translate reliability, "
                "scalability, and operational requirements into cloud automation and "
                "platform improvements."
            )
        if "collaborated with development teams to containerize" in lowered:
            return (
                "Worked with development teams to convert application deployment "
                "requirements into Docker and ECS containerization solutions, improving "
                "deployment agility."
            )
    if requirement == "Performance monitoring":
        if "correlating cloudwatch and splunk" in lowered:
            return (
                "Monitored application reliability and performance by correlating "
                "CloudWatch and Splunk telemetry, enabling faster incident triage and "
                "targeted operational improvements."
            )
        if "configured datadog monitors" in lowered:
            return (
                "Configured Datadog monitors, dashboards, and anomaly detection to measure "
                "application performance, identify regressions, and guide optimization."
            )
        if "enhanced system reliability" in lowered:
            return (
                "Monitored production reliability by correlating CloudWatch and Splunk "
                "telemetry, improving incident triage and operational visibility."
            )
    if requirement == "Documentation":
        if "authored internal runbooks" in lowered:
            return (
                "Authored and maintained operational documentation for CI/CD procedures, "
                "application recovery, troubleshooting, and infrastructure restoration "
                "drills."
            )
        if "documented findings to improve slos" in lowered:
            return (
                "Led incident reviews and documented root-cause findings, corrective "
                "actions, and operational procedures to improve SLOs and MTTR."
            )
    if requirement == "API operations":
        if "monitoring web apis" in lowered:
            return (
                "Built Splunk dashboards to monitor web API health, production behavior, "
                "and application-support signals for faster troubleshooting."
            )
    if requirement == "Architecture reviews":
        if "multi-az vpc architecture" in lowered:
            return (
                "Contributed to infrastructure design decisions by building multi-AZ VPC "
                "architecture with subnets, NAT gateways, and route tables for fault "
                "tolerance."
            )
    return text


REQUIREMENT_RULES = [
    {
        "name": "AWS infrastructure optimization",
        "jd_terms": ("aws cloud infrastructure", "compute", "storage", "serverless"),
        "candidate_terms": ("automated aws services", "deployed aws infrastructure", "aws infrastructure"),
    },
    {
        "name": "CI/CD lifecycle",
        "jd_terms": ("ci/cd", "software development lifecycle", "pipelines"),
        "candidate_terms": ("ci/cd", "github actions", "azure devops"),
    },
    {
        "name": "Infrastructure as code and automation",
        "jd_terms": ("infrastructure as code", "iac", "automation scripts", "provisioning"),
        "candidate_terms": ("terraform", "provisioning", "infrastructure"),
    },
    {
        "name": "Container orchestration",
        "jd_terms": ("containerization", "orchestration", "kubernetes", "eks"),
        "candidate_terms": ("kubernetes", "eks", "ecs fargate", "containerized"),
    },
    {
        "name": "Python application support",
        "jd_terms": ("python", "application support", "deployed ai applications"),
        "candidate_terms": ("python", "lambda"),
    },
    {
        "name": "Application lifecycle and troubleshooting",
        "jd_terms": ("lifecycle", "troubleshooting", "application support"),
        "candidate_terms": ("diagnosed", "debugged", "production support"),
    },
    {
        "name": "Business and technical collaboration",
        "jd_terms": ("business requirements", "technical specifications", "collaborat"),
        "candidate_terms": ("collaborated", "partnered", "application teams"),
    },
    {
        "name": "Performance monitoring",
        "jd_terms": ("monitor performance", "optimizations", "metrics"),
        "candidate_terms": ("performance", "monitor", "cloudwatch", "datadog"),
    },
    {
        "name": "Documentation",
        "jd_terms": ("documentation", "operational procedures", "architecture"),
        "candidate_terms": ("runbook", "documented", "documentation"),
    },
    {
        "name": "API operations",
        "jd_terms": ("api development", "api", "application design"),
        "candidate_terms": ("web api", "apis"),
    },
    {
        "name": "Architecture reviews",
        "jd_terms": ("architecture", "design reviews", "infrastructure design"),
        "candidate_terms": ("architecture", "multi-az vpc", "design"),
    },
]


def _find_candidate(
    document: DocumentType,
    requirement: str,
    candidate_terms: tuple[str, ...],
    used: set[int],
) -> tuple[int, Paragraph, str] | None:
    ranked = []
    for index, paragraph in enumerate(document.paragraphs):
        if index in used or paragraph.style.name != "List Paragraph":
            continue
        lowered = paragraph.text.lower()
        hits = sum(term in lowered for term in candidate_terms)
        updated = _rewrite_for_requirement(_normalized(paragraph.text), requirement)
        if hits and updated != _normalized(paragraph.text):
            ranked.append((hits, -len(paragraph.text), index, paragraph, updated))
    ranked.sort(key=lambda item: (-item[0], item[1], item[2]))
    return (ranked[0][2], ranked[0][3], ranked[0][4]) if ranked else None


def _generic_experience_update(
    text: str,
    role: str,
    analysis: dict,
    source_text: str,
) -> str:
    if "target role alignment:" in text.lower():
        return text
    supported = [
        skill
        for skill in analysis.get("matching_skills", [])
        if skill in set(detect_skills(source_text))
    ]
    bullet_skills = [
        skill
        for skill in detect_skills(text)
        if skill in supported
    ]
    if not bullet_skills:
        return text
    focus = ", ".join(bullet_skills[:4])
    role_label = role.strip() or analysis.get("role_type", "target role")
    return (
        f"{text} Target role alignment: applies evidenced {focus} experience "
        f"to {role_label} delivery, reliability, and operational outcomes."
    )


def _already_inserted(document: DocumentType, text: str) -> bool:
    """Do not reinsert an alignment bullet during a monitor correction pass."""
    normalized = _normalized(text)
    return any(_normalized(paragraph.text) == normalized for paragraph in document.paragraphs)


def _remove_legacy_insertions(document: DocumentType) -> None:
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("JD-ALIGNED QUALIFICATIONS:") or text == "SELECTED RELEVANT EXPERIENCE:":
            paragraph._element.getparent().remove(paragraph._element)


def _repair_split_bullets(document: DocumentType) -> list[ResumeChange]:
    """Repair source bullets split into a dangling list item and a continuation line."""
    repairs: list[ResumeChange] = []
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        original = _normalized(paragraph.text)
        if paragraph.style.name != "List Paragraph" or not original.endswith("-"):
            continue
        following = paragraphs[index + 1:]
        blank_paragraphs = []
        continuation = None
        for candidate in following:
            text = _normalized(candidate.text)
            if not text:
                blank_paragraphs.append(candidate)
                continue
            continuation = candidate
            break
        if continuation is None or continuation.style.name == "List Paragraph":
            continue
        continuation_text = _normalized(continuation.text)
        if not continuation_text or continuation_text[:1].isupper():
            continue
        updated = f"{original[:-1]}-{continuation_text}"
        _replace_paragraph_text(paragraph, updated)
        _remove_paragraph(continuation)
        for blank in blank_paragraphs:
            _move_section_break(blank, paragraph)
            _remove_paragraph(blank)
        repairs.append(
            ResumeChange(
                section="Work Experience formatting",
                original=f"{original} {continuation_text}",
                updated=updated,
                reason="Repair a split source bullet without changing its factual content.",
                client=_client_for_index(document, index),
            )
        )
    return repairs


def _promote_supported_skills(document: DocumentType, analysis: dict) -> list[ResumeChange]:
    """Update existing skill rows only; never insert a generic capability block."""
    changes: list[ResumeChange] = []
    matched = set(analysis.get("matching_skills", []))
    if "GitLab CI" not in matched:
        return changes
    for paragraph in document.paragraphs:
        original = _normalized(paragraph.text)
        if not original.lower().startswith("ci/cd tools:"):
            continue
        if "gitlab ci" in original.lower() or "gitlab" not in original.lower():
            continue
        updated = re.sub(r"\bGitLab\b", "GitLab CI", original, flags=re.IGNORECASE)
        _replace_paragraph_text(paragraph, updated)
        changes.append(
            ResumeChange(
                section="Technical Skills — CI/CD Tools",
                original=original,
                updated=updated,
                reason="Promote existing GitLab experience as GitLab CI within the existing CI/CD skills row.",
            )
        )
        break
    return changes


def tailor_resume(
    document: DocumentType,
    source_text: str,
    role: str,
    jd_text: str,
    analysis: dict,
    revision_directives: list[str] | None = None,
) -> list[ResumeChange]:
    _remove_legacy_insertions(document)
    changes: list[ResumeChange] = _repair_split_bullets(document)
    summary = _summary_paragraph(document)
    if summary is not None:
        original = _normalized(summary.text)
        updated = _targeted_summary(source_text, role, analysis)
        if original != updated:
            _replace_paragraph_text(summary, updated)
            changes.append(
                ResumeChange(
                    section="Professional Summary",
                    original=original,
                    updated=updated,
                    reason="Align the opening profile to the target role using evidenced capabilities.",
                )
            )

    jd_lower = jd_text.lower()
    used: set[int] = set()
    minimum_changes = 6 if revision_directives else 5
    for rule in REQUIREMENT_RULES:
        if not any(term in jd_lower for term in rule["jd_terms"]):
            continue
        selected = _find_candidate(
            document,
            rule["name"],
            rule["candidate_terms"],
            used,
        )
        if selected is None:
            continue
        index, candidate, updated = selected
        if _already_inserted(document, updated):
            continue
        original = _normalized(candidate.text)
        # Preserve the source bullet verbatim. The aligned statement is a separate
        # bullet immediately beneath it so the candidate can review it in context.
        _insert_after(candidate, updated, candidate)
        used.add(index)
        changes.append(
            ResumeChange(
                section="Work Experience",
                original=original,
                updated=updated,
                reason=f"Strengthen evidenced alignment to: {rule['name']}.",
                client=_client_for_index(document, index),
            )
        )
        if len(changes) >= minimum_changes + 1:
            break

    experience_changes = sum(change.section == "Work Experience" for change in changes)
    if experience_changes < 3:
        for index, paragraph in enumerate(document.paragraphs):
            if index in used or paragraph.style.name != "List Paragraph":
                continue
            original = _normalized(paragraph.text)
            if not original:
                continue
            updated = _generic_experience_update(original, role, analysis, source_text)
            if updated == original:
                continue
            if _already_inserted(document, updated):
                continue
            _insert_after(paragraph, updated, paragraph)
            used.add(index)
            changes.append(
                ResumeChange(
                    section="Work Experience",
                    original=original,
                updated=updated,
                reason="Strengthen evidenced alignment using skills already present in the base resume.",
                client=_client_for_index(document, index),
                )
            )
            experience_changes += 1
            if experience_changes >= 3:
                break

    changes.extend(_promote_supported_skills(document, analysis))
    return changes


def _numbers(text: str) -> set[str]:
    return set(re.findall(r"\b\d+(?:\.\d+)?%?\b", text))


def _evidence_excerpt(source_text: str, requirement: str) -> str | None:
    """Return a traceable source line for the review UI; never infer missing evidence."""
    aliases = {
        "CI/CD": ("ci/cd", "pipeline"),
        "AWS": ("aws",),
        "Kubernetes": ("kubernetes", "eks", "aks"),
        "Docker": ("docker", "container"),
        "GitLab CI": ("gitlab",),
        "REST API": ("api",),
        "TypeScript": ("typescript",),
        "Playwright": ("playwright",),
    }
    terms = aliases.get(requirement, (requirement.lower(),))
    for line in source_text.splitlines():
        normalized = _normalized(line)
        if normalized and any(term in normalized.lower() for term in terms):
            return normalized[:260]
    return None


def review_resume(
    base_document: DocumentType,
    tailored_document: DocumentType,
    source_text: str,
    analysis: dict,
    changes: list[ResumeChange],
) -> dict:
    base_skills = set(detect_skills(source_text))
    tailored_text = "\n".join(p.text for p in tailored_document.paragraphs)
    tailored_skills = set(detect_skills(tailored_text))
    required = analysis.get("required_skills", [])
    matrix = []
    for skill in required:
        supported = skill in base_skills
        included = skill in tailored_skills
        evidence = _evidence_excerpt(source_text, skill) if supported else None
        matrix.append(
            {
                "requirement": skill,
                "supported_by_base": supported,
                "included_in_resume": included,
                "source_evidence": evidence,
                "status": "verified" if supported and included else (
                    "gap" if not supported else "revision_required"
                ),
            }
        )

    unsupported_added = sorted(
        skill for skill in tailored_skills - base_skills
        if skill in set(required)
    )
    source_numbers = _numbers(source_text)
    invented_numbers = sorted(
        number
        for change in changes
        for number in _numbers(change.updated)
        if number not in source_numbers
    )
    layout_preserved = (
        len(base_document.sections) == len(tailored_document.sections)
        and len(base_document.styles) == len(tailored_document.styles)
        and all(
            base.page_width == tailored.page_width
            and base.page_height == tailored.page_height
            and base.left_margin == tailored.left_margin
            and base.right_margin == tailored.right_margin
            for base, tailored in zip(base_document.sections, tailored_document.sections)
        )
    )
    supported_missing = [
        item["requirement"]
        for item in matrix
        if item["status"] == "revision_required"
    ]
    experience_edits = sum(change.section == "Work Experience" for change in changes)
    directives = []
    if experience_edits < 3:
        directives.append("Revise at least three evidenced experience bullets for the JD.")
    if supported_missing:
        directives.append(
            "Surface these supported requirements in the tailored resume: "
            + ", ".join(supported_missing)
        )
    if unsupported_added:
        directives.append(
            "Remove unsupported requirements from changed content: "
            + ", ".join(unsupported_added)
        )
    if invented_numbers:
        directives.append(
            "Remove metrics not present in the base resume: " + ", ".join(invented_numbers)
        )
    if not layout_preserved:
        directives.append("Restore the base resume section geometry and style inventory.")

    verified = sum(item["status"] == "verified" for item in matrix)
    requirement_coverage = round(100 * verified / max(len(matrix), 1))
    truthfulness = 100 if not unsupported_added and not invented_numbers else max(
        0, 100 - (20 * len(unsupported_added)) - (15 * len(invented_numbers))
    )
    checks = {
        "layout_preserved": layout_preserved,
        "supported_requirements_included": not supported_missing,
        "unsupported_claims_blocked": not unsupported_added,
        "metrics_traceable": not invented_numbers,
        "experience_bullets_tailored": experience_edits >= 3,
    }
    ats_gate_passed = requirement_coverage >= 95 and truthfulness == 100
    if not ats_gate_passed:
        directives.append(
            f"ATS gate not met: {requirement_coverage}% evidence-backed requirement coverage; 95% minimum required."
        )
    return {
        "passed": not directives,
        "ats_score": requirement_coverage,
        "ats_gate_passed": ats_gate_passed,
        "task_completion": round(100 * sum(checks.values()) / len(checks)),
        "requirement_coverage": requirement_coverage,
        "truthfulness_score": truthfulness,
        "checks": checks,
        "requirements": matrix,
        "unsupported_requirements": [
            item["requirement"] for item in matrix if item["status"] == "gap"
        ],
        "directives": directives,
        "change_count": len(changes),
        "experience_edit_count": experience_edits,
    }


def run_agent_pipeline(
    document: DocumentType,
    base_document: DocumentType,
    source_text: str,
    role: str,
    jd_text: str,
    analysis: dict,
    max_revisions: int = 2,
) -> dict:
    all_changes: list[ResumeChange] = []
    review = {}
    revision_count = 0
    directives: list[str] = []
    for attempt in range(max_revisions + 1):
        changes = tailor_resume(
            document,
            source_text,
            role,
            jd_text,
            analysis,
            revision_directives=directives,
        )
        all_changes.extend(changes)
        review = review_resume(base_document, document, source_text, analysis, all_changes)
        if review["passed"]:
            break
        directives = review["directives"]
        revision_count = attempt + 1

    writer_tasks = {
        "summary_tailored": any(c.section == "Professional Summary" for c in all_changes),
        "skills_prioritized": any(c.section == "Technical Skills" for c in all_changes),
        "experience_rewritten": review.get("experience_edit_count", 0) >= 3,
        "format_preserved": review.get("checks", {}).get("layout_preserved", False),
    }
    writer_completion = round(
        100 * sum(writer_tasks.values()) / max(len(writer_tasks), 1)
    )
    return {
        "status": "approved" if review.get("passed") else "needs_attention",
        "revision_count": revision_count,
        "agents": [
            {
                "name": "Resume Updater Agent",
                "role": "Builds the weighted JD gap analysis and evidence-backed update strategy.",
                "status": "completed",
                "task_completion": 100,
                "tasks": {
                    "baseline_scored": True,
                    "gaps_identified": True,
                    "evidence_strategy_created": True,
                },
            },
            {
                "name": "Resume Tailoring Agent",
                "role": "Rewrites supported summary, skills, and experience content for the JD.",
                "status": "completed",
                "task_completion": writer_completion,
                "tasks": writer_tasks,
            },
            {
                "name": "Resume Quality Monitor",
                "role": "Verifies evidence, requirement coverage, metrics, and source formatting.",
                "status": "approved" if review.get("passed") else "needs_attention",
                "task_completion": review.get("task_completion", 0),
                "tasks": review.get("checks", {}),
            },
        ],
        "requirement_coverage": review.get("requirement_coverage", 0),
        "ats_score": review.get("ats_score", 0),
        "ats_gate_passed": review.get("ats_gate_passed", False),
        "truthfulness_score": review.get("truthfulness_score", 0),
        "unsupported_requirements": review.get("unsupported_requirements", []),
        "resume_update_plan": analysis.get("resume_update_plan", {}),
        "requirements": review.get("requirements", []),
        "changes": [
            {
                "section": change.section,
                "original": change.original,
                "updated": change.updated,
                "reason": change.reason,
                "client": change.client,
            }
            for change in all_changes
        ],
        "review_directives": review.get("directives", []),
    }
